"""Resume parsing: PDF/DOCX → plain text → ``schemas.profile.ParsedResume``.

The parser is deliberately conservative. Anything it is unsure about is reported in
``warnings`` so the user can review the result before importing it into their profile —
nothing is imported automatically.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from datetime import date
from typing import Any

from app.models import User
from app.schemas.profile import EducationIn, ExperienceIn, ParsedResume, ProjectIn, SkillIn
from app.schemas.resume import (
    ResumeBullet,
    ResumeContact,
    ResumeContent,
    ResumeEducationItem,
    ResumeExperienceItem,
    ResumeLink,
    ResumeProjectItem,
    ResumeSkillGroup,
)
from app.services import taxonomy

PDF_MAGIC = b"%PDF"
ZIP_MAGIC = b"PK\x03\x04"
PDF_MIME = "application/pdf"
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


class ResumeReadError(Exception):
    """The file could not be read (corrupt, encrypted or not a resume document)."""


# ---------------------------------------------------------------------------
# Text extraction
# ---------------------------------------------------------------------------


def detect_kind(data: bytes, filename: str) -> str | None:
    """Return "pdf" or "docx" when both the magic bytes and the extension agree."""
    ext = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    if ext == "pdf" and data.startswith(PDF_MAGIC):
        return "pdf"
    if ext == "docx" and data.startswith(ZIP_MAGIC):
        return "docx"
    return None


def _pdf_text(data: bytes) -> str:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise ResumeReadError("This PDF is password-protected. Please upload an unlocked copy.")
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    except PdfReadError as exc:
        raise ResumeReadError("We couldn't read this PDF. It may be damaged — try exporting it again.") from exc


def _docx_text(data: bytes) -> str:
    import zipfile

    from docx import Document

    try:
        doc = Document(io.BytesIO(data))
    except (zipfile.BadZipFile, KeyError, ValueError) as exc:
        raise ResumeReadError("We couldn't read this Word document. Try saving it again as .docx.") from exc
    lines: list[str] = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style = (para.style.name or "").lower() if para.style is not None else ""
        lines.append(f"• {text}" if "list" in style else text)
    for table in doc.tables:
        for row in table.rows:
            seen: list[str] = []
            for cell in row.cells:
                text = cell.text.strip()
                if text and text not in seen:  # merged cells repeat their text
                    seen.append(text)
            lines.extend(seen)
    return "\n".join(lines)


def extract_text(data: bytes, filename: str) -> str:
    kind = detect_kind(data, filename)
    if kind == "pdf":
        return _pdf_text(data)
    if kind == "docx":
        return _docx_text(data)
    raise ResumeReadError("Please upload a PDF or Word (.docx) file.")


# ---------------------------------------------------------------------------
# Patterns
# ---------------------------------------------------------------------------

EMAIL_RE = re.compile(r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+")
PHONE_RE = re.compile(r"(?:\+?1[\s.-]?)?\(?\d{3}\)?[\s.-]?\d{3}[\s.-]?\d{4}")
URL_RE = re.compile(r"(?:https?://)?(?:www\.)?[a-z0-9-]+(?:\.[a-z0-9-]+)+(?:/[^\s|,;]*)?", re.IGNORECASE)
# Bullet glyphs as they come out of PDF/DOCX text extraction (incl. Symbol-font private-use code points).
BULLET_RE = re.compile(r"^\s*(?:[•●▪■◦‣∙·\x7f\uf0b7\uf0a7\uf076\uf0d8➢►✓]\s*|[*–-]\s+|\d+[.)]\s+)")

_MONTHS = {m: i for i, m in enumerate(
    ["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"], start=1)}
_MONTH = (r"(?:jan(?:uary)?|feb(?:ruary)?|mar(?:ch)?|apr(?:il)?|may|june?|july?|aug(?:ust)?|sep(?:t(?:ember)?)?"
          r"|oct(?:ober)?|nov(?:ember)?|dec(?:ember)?)\.?")
_DATE = rf"(?:{_MONTH}\s+\d{{4}}|\d{{1,2}}/\d{{4}}|\d{{4}})"
_OPEN_END = r"(?:present|current|now|today)"
DATE_RANGE_RE = re.compile(rf"(?P<start>{_DATE})\s*(?:-|–|—|to|until)\s*(?P<end>{_DATE}|{_OPEN_END})", re.IGNORECASE)
SINGLE_DATE_RE = re.compile(
    rf"(?:expected\s+|graduat\w*\s+)?(?P<date>{_MONTH}\s+\d{{4}}|\d{{1,2}}/\d{{4}}|(?:19|20)\d{{2}})", re.IGNORECASE)
TECH_LINE_RE = re.compile(r"(?i)(tech|technologies|stack|built with)")
GPA_RE = re.compile(r"\bGPA[:\s]*([0-4](?:\.\d{1,2})?(?:\s*/\s*[0-9.]+)?)", re.IGNORECASE)

PROVINCES = {
    "AB": "Alberta", "BC": "British Columbia", "MB": "Manitoba", "NB": "New Brunswick",
    "NL": "Newfoundland and Labrador", "NS": "Nova Scotia", "NT": "Northwest Territories", "NU": "Nunavut",
    "ON": "Ontario", "PE": "Prince Edward Island", "QC": "Quebec", "SK": "Saskatchewan", "YT": "Yukon",
}
_PROVINCE_ALT = "|".join([*PROVINCES, *(re.escape(v) for v in PROVINCES.values())])
LOCATION_RE = re.compile(rf"\b([A-Z][A-Za-z.' -]{{1,40}}?),\s*({_PROVINCE_ALT})\b(?:,?\s*(Canada))?")

SECTION_ALIASES: dict[str, tuple[str, ...]] = {
    "summary": ("summary", "profile", "professional summary", "career summary", "about me", "objective",
                "career objective", "professional profile", "summary of qualifications"),
    "experience": ("experience", "work experience", "professional experience", "work history", "employment",
                   "employment history", "relevant experience", "career history"),
    "education": ("education", "education and training", "academic background", "academic history"),
    "skills": ("skills", "technical skills", "core skills", "key skills", "skills and tools", "technologies",
               "core competencies", "competencies", "skills & tools", "technical proficiencies"),
    "projects": ("projects", "personal projects", "academic projects", "selected projects", "key projects",
                 "technical projects"),
    "certifications": ("certifications", "certificates", "licenses and certifications", "licenses & certifications",
                       "certifications and training"),
    "other": ("volunteer", "volunteer experience", "volunteering", "awards", "honors", "interests", "hobbies",
              "references", "languages", "activities", "publications", "achievements", "leadership"),
}
_HEADING_INDEX = {alias: key for key, aliases in SECTION_ALIASES.items() for alias in aliases}

TITLE_WORDS = ("developer", "engineer", "analyst", "intern", "manager", "specialist", "technician", "assistant",
               "consultant", "administrator", "designer", "support", "lead", "coordinator", "associate", "architect",
               "scientist", "officer", "representative", "agent", "programmer", "tester", "co-op", "coop", "student",
               "director", "head", "clerk", "cashier", "server", "tutor", "instructor", "volunteer")
DEGREE_WORDS = ("bachelor", "master", "diploma", "certificate", "associate", "phd", "doctor", "b.sc", "bsc", "b.a",
                "ba ", "m.sc", "msc", "b.eng", "beng", "b.tech", "btech", "mba", "degree", "advanced diploma",
                "high school", "ontario secondary", "graduate certificate")
INSTITUTION_WORDS = ("university", "college", "institute", "school", "polytechnic", "academy", "cegep", "université")


def _norm_heading(line: str) -> str:
    return re.sub(r"[^a-z& ]", "", line.lower()).strip()


def heading_key(line: str) -> str | None:
    """Section key for a heading line ("WORK EXPERIENCE", "Skills:"), else None."""
    if len(line) > 45 or BULLET_RE.match(line):
        return None
    return _HEADING_INDEX.get(_norm_heading(line.rstrip(":")))


def _is_bullet(line: str) -> bool:
    return bool(BULLET_RE.match(line))


def _strip_bullet(line: str) -> str:
    return BULLET_RE.sub("", line, count=1).strip()


def parse_date(text: str | None) -> date | None:
    if not text:
        return None
    t = text.strip().lower().rstrip(".")
    if m := re.fullmatch(r"(\d{1,2})/(\d{4})", t):
        month = int(m.group(1))
        return date(int(m.group(2)), month, 1) if 1 <= month <= 12 else None
    if m := re.fullmatch(rf"({_MONTH})\s+(\d{{4}})", t, re.IGNORECASE):
        return date(int(m.group(2)), _MONTHS[m.group(1)[:3].lower()], 1)
    if re.fullmatch(r"\d{4}", t):
        return date(int(t), 1, 1)
    return None


@dataclass
class DateSpan:
    start: date | None
    end: date | None
    is_current: bool
    raw: str


def find_date_range(line: str) -> DateSpan | None:
    m = DATE_RANGE_RE.search(line)
    if not m:
        return None
    end_raw = m.group("end")
    current = bool(re.fullmatch(_OPEN_END, end_raw, re.IGNORECASE))
    return DateSpan(parse_date(m.group("start")), None if current else parse_date(end_raw), current, m.group(0))


def _clean_fragment(text: str) -> str:
    return re.sub(r"\s{2,}", " ", text).strip(" |,–—-:·•\t")


def _split_parts(line: str) -> list[str]:
    parts = re.split(r"\s+[|•·]\s+|\s+[–—-]\s+|\s+at\s+|\t|,\s+(?=[A-Z])", line)
    return [p for p in (_clean_fragment(p) for p in parts) if p]


# ---------------------------------------------------------------------------
# Contact / header
# ---------------------------------------------------------------------------


@dataclass
class _Contact:
    full_name: str | None = None
    headline: str | None = None
    email: str | None = None
    phone: str | None = None
    linkedin_url: str | None = None
    github_url: str | None = None
    portfolio_url: str | None = None
    city: str | None = None
    province: str | None = None
    country: str | None = None


def _normalize_url(url: str) -> str:
    url = url.rstrip(".,;)")
    return url if url.lower().startswith("http") else f"https://{url}"


def _looks_like_name(line: str) -> bool:
    words = line.split()
    return (2 <= len(words) <= 4 and not any(ch.isdigit() for ch in line) and "@" not in line
            and all(re.fullmatch(r"[A-Za-zÀ-ÿ'’.-]+", w) for w in words) and heading_key(line) is None
            and not any(w.lower() in TITLE_WORDS for w in words))


def parse_contact(text: str, header_lines: list[str]) -> _Contact:
    c = _Contact()
    if m := EMAIL_RE.search(text):
        c.email = m.group(0)
    if m := PHONE_RE.search(text):
        c.phone = m.group(0).strip()
    for m in URL_RE.finditer(text):
        url = m.group(0)
        low = url.lower()
        if "@" in text[max(0, m.start() - 1): m.start() + 1] or (c.email and low in c.email.lower()):
            continue
        if "linkedin.com/" in low:
            c.linkedin_url = c.linkedin_url or _normalize_url(url)
        elif "github.com/" in low:
            c.github_url = c.github_url or _normalize_url(url)
        elif (low.startswith(("http", "www.")) or re.search(r"\.(dev|io|me|ca|com|app|net|org)(/|$)", low)) \
                and not c.portfolio_url and "." in low.split("/")[0]:
            c.portfolio_url = _normalize_url(url)
    header_text = "\n".join(header_lines)
    if m := LOCATION_RE.search(header_text) or LOCATION_RE.search(text[:600]):
        c.city = _clean_fragment(m.group(1).split("|")[-1])
        prov = m.group(2)
        c.province = prov if len(prov) == 2 else next((k for k, v in PROVINCES.items() if v == prov), prov)
        c.country = "Canada"
    for line in header_lines[:4]:
        if _looks_like_name(line):
            c.full_name = line.title() if line.isupper() else line
            break
    for line in header_lines[:6]:
        if (line.lower() == (c.full_name or "").lower() or EMAIL_RE.search(line) or PHONE_RE.search(line)
                or URL_RE.fullmatch(line.strip())):
            continue
        if 1 <= len(line.split()) <= 8 and not any(ch.isdigit() for ch in line) and "," not in line:
            c.headline = line
            break
    return c


# ---------------------------------------------------------------------------
# Sections
# ---------------------------------------------------------------------------


def split_sections(lines: list[str]) -> tuple[list[str], dict[str, list[str]]]:
    """Return (header lines before the first heading, {section key: lines})."""
    header: list[str] = []
    sections: dict[str, list[str]] = {}
    current: str | None = None
    for line in lines:
        key = heading_key(line)
        if key:
            current = key
            sections.setdefault(key, [])
            continue
        # "Skills: Python, SQL" — heading and content on the same line
        if ":" in line and (inline := heading_key(line.split(":", 1)[0])) and inline in ("skills", "summary"):
            current = inline
            sections.setdefault(inline, []).append(line.split(":", 1)[1].strip())
            continue
        if current is None:
            header.append(line)
        else:
            sections[current].append(line)
    return header, sections


@dataclass
class _Entry:
    header: list[str] = field(default_factory=list)
    bullets: list[str] = field(default_factory=list)
    span: DateSpan | None = None


def _group_entries(lines: list[str]) -> list[_Entry]:
    """Group section lines into entries: header lines (title/company/dates) followed by bullets."""
    entries: list[_Entry] = []
    cur: _Entry | None = None
    for raw in lines:
        if _is_bullet(raw):
            if cur is None:
                cur = _Entry()
                entries.append(cur)
            cur.bullets.append(_strip_bullet(raw))
            continue
        line = raw.strip()
        if cur and URL_RE.fullmatch(line):
            cur.header.append(line)  # project/portfolio link under an entry
            continue
        if cur and cur.bullets and line[:1].islower():
            cur.bullets[-1] = f"{cur.bullets[-1]} {line}"  # wrapped bullet continuation
            continue
        is_sentence = len(line) > 70 or line.endswith(".")
        if cur and cur.span and is_sentence and not find_date_range(line):
            cur.bullets.append(line)  # description written without a bullet marker
            continue
        if cur is None or cur.bullets or (cur.span and find_date_range(line)) or len(cur.header) >= 3:
            cur = _Entry()
            entries.append(cur)
        span = find_date_range(line)
        if span and cur.span is None:
            cur.span = span
            line = _clean_fragment(line.replace(span.raw, " "))
        if line:
            cur.header.append(line)
    return entries


def _without_location(parts: list[str]) -> tuple[list[str], str | None]:
    location = None
    kept: list[str] = []
    for part in parts:
        if location is None and (LOCATION_RE.fullmatch(part) or part.lower() in ("remote", "hybrid")):
            location = part
        else:
            kept.append(part)
    return kept, location


def _header_parts(header: list[str]) -> tuple[list[str], str | None]:
    parts: list[str] = []
    location = None
    for line in header:
        loc = LOCATION_RE.search(line)
        if loc and location is None:
            location = _clean_fragment(loc.group(0))
            line = line.replace(loc.group(0), " ")
        parts.extend(_split_parts(line))
    kept, loc2 = _without_location(parts)
    return kept, location or loc2


def _has_title_word(text: str) -> bool:
    low = text.lower()
    return any(re.search(rf"\b{re.escape(w)}\b", low) for w in TITLE_WORDS)


def parse_experience(lines: list[str], warnings: list[str]) -> list[ExperienceIn]:
    out: list[ExperienceIn] = []
    for entry in _group_entries(lines):
        parts, location = _header_parts(entry.header)
        if not parts:
            continue
        title_idx = next((i for i, p in enumerate(parts) if _has_title_word(p)), None)
        if title_idx is None:
            title_idx = 0
            if len(parts) > 1:
                warnings.append(f"Check the job title and company for “{' | '.join(parts[:2])}”.")
        position = parts[title_idx]
        others = [p for i, p in enumerate(parts) if i != title_idx]
        company = others[0] if others else ""
        if not company:
            warnings.append(f"We couldn't find the company name for “{position}”.")
        if entry.span is None:
            warnings.append(f"We couldn't read the dates for “{position}”.")
        text = " ".join([*entry.header, *entry.bullets])
        out.append(ExperienceIn(
            company=(company or "Unknown company")[:200], position=position[:200], location=location,
            start_date=entry.span.start if entry.span else None,
            end_date=entry.span.end if entry.span else None,
            responsibilities=entry.bullets, technologies=taxonomy.extract_skills(text),
            sort_order=len(out),
        ))
    return out


def _split_degree(text: str) -> tuple[str | None, str | None]:
    """"Bachelor of Science in Computer Science" -> ("Bachelor of Science", "Computer Science")."""
    m = re.match(r"(.+?)\s+(?:in|of|,|-|–)\s+(.+)$", text)
    if m and re.search(r"\bof\b", m.group(1)) is None and " of " in text:
        m = re.match(r"(.+?\bof\s+\w+(?:\s+\w+)?)\s+(?:in|,|-|–)\s+(.+)$", text) or m
    if m and any(w in m.group(1).lower() for w in DEGREE_WORDS):
        return m.group(1).strip(), m.group(2).strip()
    return text.strip(), None


def parse_education(lines: list[str], warnings: list[str]) -> list[EducationIn]:
    items: list[dict[str, Any]] = []
    cur: dict[str, Any] | None = None
    for raw in lines:
        line = _strip_bullet(raw) if _is_bullet(raw) else raw.strip()
        low = line.lower()
        is_institution = any(w in low for w in INSTITUTION_WORDS)
        if cur is None or (is_institution and cur.get("institution")):
            cur = {"coursework": []}
            items.append(cur)
        if low.startswith(("relevant coursework", "coursework", "courses")):
            courses = line.split(":", 1)[-1]
            cur["coursework"].extend(c.strip() for c in re.split(r"[,;]", courses) if c.strip())
            continue
        if gpa := GPA_RE.search(line):
            cur["gpa"] = gpa.group(1).strip()
            line = _clean_fragment(line.replace(gpa.group(0), " "))
        span = find_date_range(line)
        if span:
            cur["start_date"], cur["end_date"] = span.start, span.end
            line = _clean_fragment(line.replace(span.raw, " "))
        elif single := SINGLE_DATE_RE.search(line):
            cur["end_date"] = parse_date(single.group("date"))
            line = _clean_fragment(line.replace(single.group(0), " "))
        loc = LOCATION_RE.search(line)
        if loc:
            cur["location"] = _clean_fragment(loc.group(0))
            line = _clean_fragment(line.replace(loc.group(0), " "))
        for part in _split_parts(line) if line else []:
            plow = part.lower()
            if any(w in plow for w in INSTITUTION_WORDS) and not cur.get("institution"):
                cur["institution"] = part
            elif any(w in f"{plow} " for w in DEGREE_WORDS) and not cur.get("degree"):
                cur["degree"], program = _split_degree(part)
                cur["program"] = cur.get("program") or program
            elif not cur.get("program") and cur.get("degree"):
                cur["program"] = part
    result: list[EducationIn] = []
    for i, item in enumerate(items):
        institution = item.get("institution") or item.get("degree")
        if not institution:
            continue
        if not item.get("institution"):
            warnings.append(f"We couldn't find the school name for “{institution}”.")
        result.append(EducationIn(
            institution=institution[:200], degree=item.get("degree") if item.get("institution") else None,
            program=item.get("program"), start_date=item.get("start_date"), end_date=item.get("end_date"),
            gpa=item.get("gpa"), location=item.get("location"), coursework=item["coursework"], sort_order=i,
        ))
    return result


def parse_skills(lines: list[str]) -> list[SkillIn]:
    seen: set[str] = set()
    skills: list[SkillIn] = []

    def add(name: str) -> None:
        canonical = taxonomy.canonicalize(name)
        if canonical.lower() in seen or not canonical:
            return
        seen.add(canonical.lower())
        skills.append(SkillIn(name=canonical[:120], category=taxonomy.category_of(canonical)))  # type: ignore[arg-type]

    for raw in lines:
        line = _strip_bullet(raw) if _is_bullet(raw) else raw
        if ":" in line:
            line = line.split(":", 1)[1]
        for token in re.split(r"[,|;•·/]|\s{2,}|\band\b", line):
            token = _clean_fragment(token.strip(" ()"))
            if token and len(token) <= 40 and len(token.split()) <= 4 and not token[0].isdigit():
                add(token)
        for found in taxonomy.extract_skills(line):
            add(found)
    return skills


def parse_projects(lines: list[str]) -> list[ProjectIn]:
    projects: list[ProjectIn] = []
    for entry in _group_entries(lines):
        if not entry.header and not entry.bullets:
            continue
        head = entry.header[0] if entry.header else entry.bullets.pop(0)
        link = next((h.strip() for h in entry.header if URL_RE.fullmatch(h.strip())), None)
        tech_line = next((h for h in entry.header[1:] if TECH_LINE_RE.match(h)), None)
        name_parts = [p for p in re.split(r"\s+[|–—-]\s+|:\s+", head) if p.strip()]
        name = _clean_fragment(name_parts[0]) or head
        description = " ".join(_clean_fragment(p) for p in name_parts[1:]) or None
        extra = [h for h in entry.header[1:] if h != tech_line and not URL_RE.fullmatch(h.strip())]
        if extra and not description:
            description = " ".join(extra)
        text = " ".join([*entry.header, *entry.bullets])
        url = _normalize_url(link) if link else None
        projects.append(ProjectIn(
            name=name[:200], description=description, technologies=taxonomy.extract_skills(text),
            responsibilities=entry.bullets,
            github_url=url if url and "github.com" in url else None,
            demo_url=url if url and "github.com" not in url else None, sort_order=len(projects),
        ))
    return projects


def parse_certifications(lines: list[str]) -> list[SkillIn]:
    certs: list[SkillIn] = []
    for raw in lines:
        line = _strip_bullet(raw) if _is_bullet(raw) else raw
        line = _clean_fragment(re.sub(r"\(\s*\)", "", DATE_RANGE_RE.sub("", SINGLE_DATE_RE.sub("", line))))
        if line:
            certs.append(SkillIn(name=line[:120], category="certifications"))  # keep the exact credential name
    return certs


def _lines(text: str) -> list[str]:
    out = []
    for raw in text.replace("\r", "\n").split("\n"):
        line = re.sub(r"[ \t ]+", " ", raw).strip()
        if line:
            out.append(line)
    return out


def parse_text(text: str) -> ParsedResume:
    lines = _lines(text)
    warnings: list[str] = []
    if len(text.strip()) < 80:
        warnings.append("Very little text could be read. If this is a scanned PDF, try uploading a Word version.")
    header, sections = split_sections(lines)
    contact = parse_contact(text, header or lines[:6])
    if not contact.full_name:
        warnings.append("We couldn't find your name at the top of the resume.")
    if not contact.email:
        warnings.append("No email address was found.")

    skills = parse_skills(sections.get("skills", []))
    if not skills:
        skills = [SkillIn(name=s, category=taxonomy.category_of(s))  # type: ignore[arg-type]
                  for s in taxonomy.extract_skills(text)]
        if skills:
            warnings.append("No skills section was found; skills were detected from the rest of your resume.")
    skills.extend(c for c in parse_certifications(sections.get("certifications", []))
                  if c.name.lower() not in {s.name.lower() for s in skills})

    experiences = parse_experience(sections.get("experience", []), warnings)
    if not experiences:
        warnings.append("No work experience section was detected.")
    educations = parse_education(sections.get("education", []), warnings)
    summary = " ".join(sections.get("summary", [])) or None

    return ParsedResume(
        full_name=contact.full_name, email=contact.email, phone=contact.phone, city=contact.city,
        province=contact.province, country=contact.country, linkedin_url=contact.linkedin_url,
        github_url=contact.github_url, portfolio_url=contact.portfolio_url, headline=contact.headline,
        summary=summary, skills=skills, experiences=experiences, educations=educations,
        projects=parse_projects(sections.get("projects", [])), warnings=warnings,
    )


def parse_file(data: bytes, filename: str) -> tuple[str, ParsedResume]:
    """Extract text and parse it. Raises ``ResumeReadError`` for unreadable files."""
    text = extract_text(data, filename)
    return text, parse_text(text)


# ---------------------------------------------------------------------------
# Parsed resume -> structured content
# ---------------------------------------------------------------------------


def _display(d: date | None) -> str | None:
    return d.strftime("%b %Y") if d else None


def parsed_to_content(parsed: ParsedResume, user: User) -> dict[str, Any]:
    """Structured ``ResumeContent`` for an uploaded resume (so it can be tailored and rendered)."""
    links = [ResumeLink(label=label, url=url) for label, url in (
        ("LinkedIn", parsed.linkedin_url), ("GitHub", parsed.github_url), ("Portfolio", parsed.portfolio_url)) if url]
    location = ", ".join(p for p in (parsed.city, parsed.province) if p) or None
    groups: dict[str, list[str]] = {}
    for s in parsed.skills:
        if s.category != "certifications":
            groups.setdefault(taxonomy.CATEGORY_LABELS.get(s.category or "other", "Other"), []).append(s.name)
    content = ResumeContent(
        contact=ResumeContact(name=parsed.full_name or user.full_name, email=parsed.email or user.email,
                              phone=parsed.phone, location=location, links=links),
        headline=parsed.headline,
        summary=parsed.summary,
        skills=[ResumeSkillGroup(category=cat, items=items) for cat, items in groups.items()],
        experience=[
            ResumeExperienceItem(
                id=f"exp-p{i}", company=e.company, position=e.position, location=e.location,
                start=_display(e.start_date), end=_display(e.end_date) or ("Present" if e.start_date else None),
                bullets=[ResumeBullet(id=f"exp-p{i}-b{j}", text=t) for j, t in enumerate(e.responsibilities, 1)],
                technologies=e.technologies,
            ) for i, e in enumerate(parsed.experiences, 1)
        ],
        projects=[
            ResumeProjectItem(
                id=f"proj-p{i}", name=p.name, description=p.description, technologies=p.technologies,
                bullets=[ResumeBullet(id=f"proj-p{i}-b{j}", text=t) for j, t in enumerate(p.responsibilities, 1)],
                url=p.github_url or p.demo_url,
            ) for i, p in enumerate(parsed.projects, 1)
        ],
        education=[
            ResumeEducationItem(
                id=f"edu-p{i}", institution=ed.institution, degree=ed.degree, program=ed.program,
                start=_display(ed.start_date), end=_display(ed.end_date), gpa=ed.gpa, location=ed.location,
                details=[f"Relevant coursework: {', '.join(ed.coursework)}"] if ed.coursework else [],
            ) for i, ed in enumerate(parsed.educations, 1)
        ],
        certifications=[s.name for s in parsed.skills if s.category == "certifications"],
    )
    return content.model_dump()
