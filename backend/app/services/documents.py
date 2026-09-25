"""Render resumes and cover letters to PDF (reportlab) and DOCX (python-docx).

Layout goals: single column, ATS-friendly (real text, standard fonts, no images or text
boxes), clear section headings, consistent spacing and clickable links.
"""

from __future__ import annotations

import io
from datetime import date
from typing import Any, Literal
from xml.sax.saxutils import escape

from app.schemas.resume import ResumeContact, ResumeContent

Format = Literal["pdf", "docx"]
PDF_MIME = "application/pdf"
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
MIME_TYPES: dict[str, str] = {"pdf": PDF_MIME, "docx": DOCX_MIME}

INK = "#16181d"
MUTED = "#555b66"
ACCENT = "#1f3a5f"

SECTION_TITLES = {"summary": "Summary", "skills": "Skills", "experience": "Experience", "projects": "Projects",
                  "education": "Education", "certifications": "Certifications"}


def with_extension(file_name: str, fmt: Format) -> str:
    stem = file_name.rsplit(".", 1)[0] if "." in file_name else file_name
    return f"{stem}.{fmt}"


def _date_range(start: str | None, end: str | None) -> str:
    if start and end:
        return f"{start} – {end}"
    return start or end or ""


def _contact_parts(contact: ResumeContact) -> list[tuple[str, str | None]]:
    """(text, url) pairs for the contact line."""
    parts: list[tuple[str, str | None]] = []
    if contact.email:
        parts.append((contact.email, f"mailto:{contact.email}"))
    if contact.phone:
        parts.append((contact.phone, None))
    if contact.location:
        parts.append((contact.location, None))
    parts += [(link.url.removeprefix("https://").removeprefix("http://").removeprefix("www."), link.url)
              for link in contact.links]
    return parts


def _education_title(degree: str | None, program: str | None) -> str:
    if degree and program:
        return f"{degree}, {program}"
    return degree or program or ""


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------


def _pdf_styles() -> dict[str, Any]:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_RIGHT
    from reportlab.lib.styles import ParagraphStyle

    ink, muted, accent = colors.HexColor(INK), colors.HexColor(MUTED), colors.HexColor(ACCENT)
    base = ParagraphStyle("base", fontName="Helvetica", fontSize=9.8, leading=13.2, textColor=ink)
    return {
        "name": ParagraphStyle("name", parent=base, fontName="Helvetica-Bold", fontSize=20, leading=24),
        "headline": ParagraphStyle("headline", parent=base, fontSize=11, leading=14, textColor=muted),
        "contact": ParagraphStyle("contact", parent=base, fontSize=8.8, leading=12, textColor=muted),
        "section": ParagraphStyle("section", parent=base, fontName="Helvetica-Bold", fontSize=10.5, leading=13,
                                  textColor=accent, spaceBefore=10, spaceAfter=2),
        "body": base,
        "item": ParagraphStyle("item", parent=base, fontName="Helvetica-Bold", fontSize=10),
        "sub": ParagraphStyle("sub", parent=base, textColor=muted, fontSize=9.2),
        "date": ParagraphStyle("date", parent=base, textColor=muted, fontSize=9.2, alignment=TA_RIGHT),
        "bullet": ParagraphStyle("bullet", parent=base, leftIndent=12, bulletIndent=2, spaceBefore=1.2),
        "letter": ParagraphStyle("letter", parent=base, fontSize=10.5, leading=15, spaceAfter=9),
    }


def _p(text: str, style: Any, bullet: str | None = None) -> Any:
    from reportlab.platypus import Paragraph

    return Paragraph(text, style, bulletText=bullet)


def _row(left: str, right: str, styles: dict[str, Any], width: float) -> Any:
    """Title on the left, dates right-aligned — a borderless two-column row."""
    from reportlab.platypus import Table, TableStyle

    table = Table([[_p(left, styles["item"]), _p(escape(right), styles["date"])]],
                  colWidths=[width * 0.72, width * 0.28])
    table.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "BOTTOM"), ("LEFTPADDING", (0, 0), (-1, -1), 0),
                               ("RIGHTPADDING", (0, 0), (-1, -1), 0), ("TOPPADDING", (0, 0), (-1, -1), 5),
                               ("BOTTOMPADDING", (0, 0), (-1, -1), 0)]))
    return table


def _pdf_section(key: str, c: ResumeContent, styles: dict[str, Any], width: float) -> list[Any]:
    from reportlab.lib import colors
    from reportlab.platypus import HRFlowable

    flow: list[Any] = []

    def bullets(items: list[str]) -> None:
        flow.extend(_p(escape(t), styles["bullet"], "•") for t in items)

    if key == "summary" and c.summary:
        flow.append(_p(escape(c.summary), styles["body"]))
    elif key == "skills" and c.skills:
        flow += [_p(f"<b>{escape(g.category)}:</b> {escape(', '.join(g.items))}", styles["body"])
                 for g in c.skills if g.items]
    elif key == "experience" and c.experience:
        for e in c.experience:
            flow.append(_row(f"{escape(e.position)}", _date_range(e.start, e.end), styles, width))
            flow.append(_p(escape(" · ".join(x for x in (e.company, e.location) if x)), styles["sub"]))
            bullets([b.text for b in e.bullets])
    elif key == "projects" and c.projects:
        for pr in c.projects:
            title = escape(pr.name)
            if pr.url:
                title += f' <font size="8.8" color="{MUTED}"><a href="{escape(pr.url)}">{escape(pr.url)}</a></font>'
            flow.append(_row(title, "", styles, width))
            meta = " · ".join(x for x in (pr.description, ", ".join(pr.technologies)) if x)
            if meta:
                flow.append(_p(escape(meta), styles["sub"]))
            bullets([b.text for b in pr.bullets])
    elif key == "education" and c.education:
        for ed in c.education:
            flow.append(_row(escape(_education_title(ed.degree, ed.program) or ed.institution),
                             _date_range(ed.start, ed.end), styles, width))
            sub = " · ".join(x for x in (ed.institution if (ed.degree or ed.program) else None, ed.location,
                                          f"GPA {ed.gpa}" if ed.gpa else None) if x)
            if sub:
                flow.append(_p(escape(sub), styles["sub"]))
            bullets(ed.details)
    elif key == "certifications" and c.certifications:
        bullets(c.certifications)
    if not flow:
        return []
    heading = [_p(SECTION_TITLES[key].upper(), styles["section"]),
               HRFlowable(width="100%", thickness=0.6, color=colors.HexColor("#c9ced6"), spaceBefore=0, spaceAfter=4)]
    return heading + flow


def _pdf_doc(buffer: io.BytesIO, title: str, author: str) -> Any:
    from reportlab.lib.pagesizes import LETTER
    from reportlab.lib.units import inch
    from reportlab.platypus import SimpleDocTemplate

    return SimpleDocTemplate(buffer, pagesize=LETTER, leftMargin=0.75 * inch, rightMargin=0.75 * inch,
                             topMargin=0.6 * inch, bottomMargin=0.6 * inch, title=title, author=author,
                             creator="Applier")


def _pdf_header(contact: ResumeContact, headline: str | None, styles: dict[str, Any]) -> list[Any]:
    from reportlab.platypus import Spacer

    flow: list[Any] = [_p(escape(contact.name or ""), styles["name"])]
    if headline:
        flow.append(_p(escape(headline), styles["headline"]))
    parts = [f'<a href="{escape(url)}">{escape(text)}</a>' if url else escape(text)
             for text, url in _contact_parts(contact)]
    if parts:
        flow.append(_p("  ·  ".join(parts), styles["contact"]))
    flow.append(Spacer(1, 4))
    return flow


def resume_pdf(content: dict[str, Any]) -> bytes:
    c = ResumeContent.model_validate(content)
    styles = _pdf_styles()
    buffer = io.BytesIO()
    doc = _pdf_doc(buffer, f"{c.contact.name} — Resume", c.contact.name)
    flow = _pdf_header(c.contact, c.headline, styles)
    for key in c.section_order:
        if key in SECTION_TITLES:
            flow += _pdf_section(key, c, styles, doc.width)
    doc.build(flow)
    return buffer.getvalue()


def _letter_paragraphs(text: str) -> list[str]:
    return [p.strip() for p in text.replace("\r\n", "\n").split("\n\n") if p.strip()]


def cover_letter_pdf(text: str, contact: ResumeContact, *, company: str | None = None,
                     on: date | None = None) -> bytes:
    from reportlab.platypus import Spacer

    styles = _pdf_styles()
    buffer = io.BytesIO()
    doc = _pdf_doc(buffer, f"{contact.name} — Cover Letter", contact.name)
    flow = _pdf_header(contact, None, styles)
    flow.append(Spacer(1, 10))
    flow.append(_p(escape((on or date.today()).strftime("%B %-d, %Y")), styles["letter"]))
    if company:
        flow.append(_p(f"Hiring Team<br/>{escape(company)}", styles["letter"]))
    flow += [_p(escape(para).replace("\n", "<br/>"), styles["letter"]) for para in _letter_paragraphs(text)]
    doc.build(flow)
    return buffer.getvalue()


# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------


def _docx_base() -> Any:
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor

    doc = Document()
    for section in doc.sections:
        section.left_margin = section.right_margin = Inches(0.75)
        section.top_margin = section.bottom_margin = Inches(0.6)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK[1:])
    normal.paragraph_format.space_after = Pt(2)
    return doc


def _bottom_border(paragraph: Any) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for attr, value in (("w:val", "single"), ("w:sz", "4"), ("w:space", "1"), ("w:color", "C9CED6")):
        bottom.set(qn(attr), value)
    borders.append(bottom)
    p_pr.append(borders)


def _docx_run(paragraph: Any, text: str, *, bold: bool = False, size: float | None = None,
              color: str | None = None) -> Any:
    from docx.shared import Pt, RGBColor

    run = paragraph.add_run(text)
    run.bold = bold
    if size:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color[1:])
    return run


def _docx_header(doc: Any, contact: ResumeContact, headline: str | None) -> None:
    _docx_run(doc.add_paragraph(), contact.name or "", bold=True, size=20)
    if headline:
        _docx_run(doc.add_paragraph(), headline, size=11.5, color=MUTED)
    parts = [text for text, _ in _contact_parts(contact)]
    if parts:
        _docx_run(doc.add_paragraph(), "  ·  ".join(parts), size=9.5, color=MUTED)


def _docx_row(doc: Any, left: str, right: str) -> None:
    from docx.enum.text import WD_TAB_ALIGNMENT
    from docx.shared import Pt

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    width = doc.sections[0].page_width - doc.sections[0].left_margin - doc.sections[0].right_margin
    p.paragraph_format.tab_stops.add_tab_stop(width, WD_TAB_ALIGNMENT.RIGHT)
    _docx_run(p, left, bold=True)
    if right:
        _docx_run(p, f"\t{right}", color=MUTED, size=9.5)


def _docx_sub(doc: Any, text: str) -> None:
    if text:
        _docx_run(doc.add_paragraph(), text, color=MUTED, size=9.5)


def _docx_bullets(doc: Any, items: list[str]) -> None:
    for text in items:
        doc.add_paragraph(text, style="List Bullet")


def _docx_section(doc: Any, key: str, c: ResumeContent) -> None:
    from docx.shared import Pt

    has_content = {"summary": bool(c.summary), "skills": bool(c.skills), "experience": bool(c.experience),
                   "projects": bool(c.projects), "education": bool(c.education),
                   "certifications": bool(c.certifications)}
    if not has_content.get(key):
        return
    heading = doc.add_paragraph()
    heading.paragraph_format.space_before = Pt(10)
    _docx_run(heading, SECTION_TITLES[key].upper(), bold=True, size=10.5, color=ACCENT)
    _bottom_border(heading)
    if key == "summary":
        doc.add_paragraph(c.summary)
    elif key == "skills":
        for g in c.skills:
            p = doc.add_paragraph()
            _docx_run(p, f"{g.category}: ", bold=True)
            _docx_run(p, ", ".join(g.items))
    elif key == "experience":
        for e in c.experience:
            _docx_row(doc, e.position, _date_range(e.start, e.end))
            _docx_sub(doc, " · ".join(x for x in (e.company, e.location) if x))
            _docx_bullets(doc, [b.text for b in e.bullets])
    elif key == "projects":
        for pr in c.projects:
            _docx_row(doc, pr.name, pr.url or "")
            _docx_sub(doc, " · ".join(x for x in (pr.description, ", ".join(pr.technologies)) if x))
            _docx_bullets(doc, [b.text for b in pr.bullets])
    elif key == "education":
        for ed in c.education:
            _docx_row(doc, _education_title(ed.degree, ed.program) or ed.institution, _date_range(ed.start, ed.end))
            _docx_sub(doc, " · ".join(x for x in (ed.institution if (ed.degree or ed.program) else None, ed.location,
                                                  f"GPA {ed.gpa}" if ed.gpa else None) if x))
            _docx_bullets(doc, ed.details)
    elif key == "certifications":
        _docx_bullets(doc, c.certifications)


def _docx_bytes(doc: Any, title: str, author: str) -> bytes:
    doc.core_properties.title = title
    doc.core_properties.author = author
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def resume_docx(content: dict[str, Any]) -> bytes:
    c = ResumeContent.model_validate(content)
    doc = _docx_base()
    _docx_header(doc, c.contact, c.headline)
    for key in c.section_order:
        if key in SECTION_TITLES:
            _docx_section(doc, key, c)
    return _docx_bytes(doc, f"{c.contact.name} — Resume", c.contact.name)


def cover_letter_docx(text: str, contact: ResumeContact, *, company: str | None = None,
                      on: date | None = None) -> bytes:
    from docx.shared import Pt

    doc = _docx_base()
    _docx_header(doc, contact, None)
    doc.add_paragraph()
    doc.add_paragraph((on or date.today()).strftime("%B %-d, %Y"))
    if company:
        doc.add_paragraph(f"Hiring Team\n{company}")
    for para in _letter_paragraphs(text):
        p = doc.add_paragraph(para)
        p.paragraph_format.space_after = Pt(9)
    return _docx_bytes(doc, f"{contact.name} — Cover Letter", contact.name)


# ---------------------------------------------------------------------------
# Public API
# ---------------------------------------------------------------------------


def render_resume(content: dict[str, Any], fmt: Format = "pdf") -> bytes:
    return resume_pdf(content) if fmt == "pdf" else resume_docx(content)


def render_cover_letter(text: str, contact: ResumeContact, fmt: Format = "pdf", *, company: str | None = None,
                        on: date | None = None) -> bytes:
    if fmt == "pdf":
        return cover_letter_pdf(text, contact, company=company, on=on)
    return cover_letter_docx(text, contact, company=company, on=on)
