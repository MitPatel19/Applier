"""Resume parsing from real PDF/DOCX files, plus the upload API."""

from __future__ import annotations

import io
from datetime import date

import pytest

from app.schemas.resume import ResumeContent
from app.services import resume_parser

RESUME_LINES = [
    "Mit Patel",
    "Junior Software Developer",
    "Thunder Bay, ON | mit.patel@example.com | (807) 555-0142",
    "linkedin.com/in/mitpatel | github.com/mitpatel | mitpatel.dev",
    "SUMMARY",
    "Developer with two years building web apps in Python and React.",
    "EXPERIENCE",
    "Software Developer Intern | Northern Tech Solutions | Thunder Bay, ON",
    "Jan 2022 – Present",
    "• Built REST APIs with Django and Postgres for 3 internal tools",
    "• Maintained the CI pipeline in GitHub Actions",
    "Web Developer, Acme Corp",
    "2020 - 2021",
    "• Worked on the customer portal using React and TypeScript",
    "IT Support Assistant at City Library 05/2019 - 08/2019",
    "• Resolved about 30 tickets a week",
    "EDUCATION",
    "Lakehead University",
    "Bachelor of Science in Computer Science, 2018 - 2022",
    "GPA: 3.6/4.0",
    "Relevant coursework: Data Structures, Databases",
    "TECHNICAL SKILLS",
    "Languages: Python, JavaScript, TypeScript, SQL",
    "Tools: Git, Docker, Jira",
    "PROJECTS",
    "Budget Tracker | Personal finance web app",
    "• Built with Flask and SQLite",
    "github.com/mitpatel/budget",
    "CERTIFICATIONS",
    "AWS Cloud Practitioner (2023)",
]


def make_pdf(lines: list[str]) -> bytes:
    from reportlab.lib.pagesizes import LETTER
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfgen import canvas

    buffer = io.BytesIO()
    c = canvas.Canvas(buffer, pagesize=LETTER)
    y = 750
    for line in lines:
        c.setFont("Helvetica", 10)
        c.drawString(50, y, line if pdfmetrics.stringWidth(line, "Helvetica", 10) < 520 else line[:90])
        y -= 16
    c.save()
    return buffer.getvalue()


def make_docx() -> bytes:
    from docx import Document

    doc = Document()
    doc.add_paragraph("Priya Sharma")
    doc.add_paragraph("priya.sharma@example.com · 416-555-0199 · Toronto, Ontario")
    doc.add_paragraph("Work Experience")
    doc.add_paragraph("Data Analyst — Maple Insights")
    doc.add_paragraph("05/2020 - 08/2021")
    doc.add_paragraph("Built Excel and SQL reports for the finance team", style="List Bullet")
    doc.add_paragraph("Education")
    doc.add_paragraph("Diploma, Computer Programming — Seneca College 2018 - 2020")
    doc.add_paragraph("Skills")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Python, Pandas"
    table.rows[0].cells[1].text = "Tableau, Excel"
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def test_parses_pdf_resume():
    text, parsed = resume_parser.parse_file(make_pdf(RESUME_LINES), "resume.pdf")
    assert "Northern Tech Solutions" in text
    assert parsed.full_name == "Mit Patel"
    assert parsed.headline == "Junior Software Developer"
    assert parsed.email == "mit.patel@example.com" and parsed.phone == "(807) 555-0142"
    assert parsed.linkedin_url == "https://linkedin.com/in/mitpatel"
    assert parsed.github_url == "https://github.com/mitpatel"
    assert parsed.portfolio_url == "https://mitpatel.dev"
    assert (parsed.city, parsed.province, parsed.country) == ("Thunder Bay", "ON", "Canada")
    assert parsed.summary.startswith("Developer with two years")

    intern, web, support = parsed.experiences
    assert (intern.position, intern.company, intern.location) == ("Software Developer Intern",
                                                                  "Northern Tech Solutions", "Thunder Bay, ON")
    assert intern.start_date == date(2022, 1, 1) and intern.end_date is None
    assert intern.responsibilities == ["Built REST APIs with Django and Postgres for 3 internal tools",
                                       "Maintained the CI pipeline in GitHub Actions"]
    assert {"Django", "PostgreSQL", "GitHub Actions"} <= set(intern.technologies)
    assert (web.position, web.company, web.start_date, web.end_date) == ("Web Developer", "Acme Corp",
                                                                         date(2020, 1, 1), date(2021, 1, 1))
    assert (support.position, support.company) == ("IT Support Assistant", "City Library")
    assert (support.start_date, support.end_date) == (date(2019, 5, 1), date(2019, 8, 1))

    (edu,) = parsed.educations
    assert (edu.institution, edu.degree, edu.program, edu.gpa) == ("Lakehead University", "Bachelor of Science",
                                                                   "Computer Science", "3.6/4.0")
    assert edu.coursework == ["Data Structures", "Databases"]
    skills = {s.name: s.category for s in parsed.skills}
    assert skills["Python"] == "programming" and skills["Docker"] == "devops"
    assert skills["AWS Cloud Practitioner"] == "certifications"
    (project,) = parsed.projects
    assert project.name == "Budget Tracker" and project.github_url == "https://github.com/mitpatel/budget"
    assert {"Flask", "SQLite"} <= set(project.technologies)


def test_parses_docx_with_list_styles_and_tables():
    text, parsed = resume_parser.parse_file(make_docx(), "priya.docx")
    assert parsed.full_name == "Priya Sharma" and parsed.email == "priya.sharma@example.com"
    assert (parsed.city, parsed.province) == ("Toronto", "ON")
    (exp,) = parsed.experiences
    assert (exp.position, exp.company) == ("Data Analyst", "Maple Insights")
    assert (exp.start_date, exp.end_date) == (date(2020, 5, 1), date(2021, 8, 1))
    assert exp.responsibilities == ["Built Excel and SQL reports for the finance team"]
    (edu,) = parsed.educations
    assert edu.institution == "Seneca College" and edu.degree == "Diploma"
    assert {"Python", "Pandas", "Excel"} <= {s.name for s in parsed.skills}


def test_uncertain_parts_become_warnings():
    parsed = resume_parser.parse_text("EXPERIENCE\nSomething Somewhere\n• did things\n")
    assert any("name" in w for w in parsed.warnings)
    assert any("email" in w.lower() for w in parsed.warnings)
    assert any("dates" in w for w in parsed.warnings)


@pytest.mark.parametrize("line,start,end", [
    ("Jan 2022 – Present", date(2022, 1, 1), None),
    ("2021 - 2023", date(2021, 1, 1), date(2023, 1, 1)),
    ("05/2020 - 08/2021", date(2020, 5, 1), date(2021, 8, 1)),
    ("September 2019 to December 2020", date(2019, 9, 1), date(2020, 12, 1)),
])
def test_date_ranges(line, start, end):
    span = resume_parser.find_date_range(line)
    assert span is not None and (span.start, span.end) == (start, end)


def test_rejects_mismatched_file_types():
    pdf = make_pdf(["Hello"])
    assert resume_parser.detect_kind(pdf, "resume.pdf") == "pdf"
    assert resume_parser.detect_kind(pdf, "resume.docx") is None
    assert resume_parser.detect_kind(b"just text", "resume.pdf") is None


def test_parsed_to_content_is_valid(user):
    _, parsed = resume_parser.parse_file(make_pdf(RESUME_LINES), "resume.pdf")
    content = ResumeContent.model_validate(resume_parser.parsed_to_content(parsed, user))
    assert content.contact.name == "Mit Patel" and content.experience[0].id == "exp-p1"
    assert content.experience[0].bullets[0].id == "exp-p1-b1"
    assert content.certifications == ["AWS Cloud Practitioner"]


def test_upload_api(auth_client):
    pdf = make_pdf(RESUME_LINES)
    r = auth_client.post("/api/resumes", data={"name": "Main resume"},
                         files={"file": ("Mit Resume.pdf", pdf, "application/pdf")})
    assert r.status_code == 201, r.text
    resume = r.json()
    assert resume["is_default"] and resume["has_file"] and resume["skills_count"] >= 5
    assert resume["content"]["experience"][0]["company"] == "Northern Tech Solutions"

    parsed = auth_client.get(f"/api/resumes/{resume['id']}/parsed").json()
    assert parsed["full_name"] == "Mit Patel"
    original = auth_client.get(f"/api/resumes/{resume['id']}/file")
    assert original.content == pdf and "attachment" in original.headers["content-disposition"]

    r = auth_client.post("/api/resumes", data={"name": "Second", "set_default": "true"},
                         files={"file": ("second.pdf", pdf, "application/pdf")})
    listed = auth_client.get("/api/resumes").json()
    assert [(x["id"], x["is_default"]) for x in listed] == [(r.json()["id"], True), (resume["id"], False)]

    bad = auth_client.post("/api/resumes", data={"name": "Bad"}, files={"file": ("notes.txt", b"hello", "text/plain")})
    assert bad.status_code == 415 and bad.json()["error"]["message"] == "Please upload a PDF or Word (.docx) file."

    r = auth_client.patch(f"/api/resumes/{resume['id']}", json={"content": {**resume["content"], "summary": "New"}})
    assert r.json()["version"] == 2
    assert auth_client.get(f"/api/resumes/{resume['id']}/render?format=docx").content.startswith(b"PK")
    assert auth_client.delete(f"/api/resumes/{r.json()['id']}").status_code == 204
